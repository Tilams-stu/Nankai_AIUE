from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_ROOT = ROOT / "outputs" / "03_日报周报"


def set_run_font(run, font_name="宋体", size=11, bold=False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_style(paragraph, before=0, after=6, line=1.15):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def set_body_style(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15


def add_title(doc, text):
    p = doc.add_paragraph()
    set_paragraph_style(p, after=10, line=1.15)
    r = p.add_run(text)
    set_run_font(r, "微软雅黑", 16, True, "1F4D78")
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    set_paragraph_style(p, before=8, after=4, line=1.15)
    r = p.add_run(text)
    set_run_font(r, "微软雅黑", 12, True, "2E74B5")
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph_style(p, after=6, line=1.15)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_footer_page_number(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = 2
    run = p.add_run("第 ")
    set_run_font(run, "宋体", 9, False, "666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    run2 = p.add_run(" 页")
    set_run_font(run2, "宋体", 9, False, "666666")


def build_report(day, title, module, intro, sections):
    out_dir = OUT_ROOT / day
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"日报{day}.docx"

    doc = Document()
    set_body_style(doc)
    add_title(doc, title)
    add_para(doc, f"负责模块：{module}", bold_prefix="负责模块：")
    add_para(doc, intro)

    for heading, paragraphs in sections:
        add_heading(doc, heading)
        for text in paragraphs:
            add_para(doc, text)

    add_footer_page_number(doc)
    doc.save(out_path)
    return out_path


REPORTS = [
    {
        "day": "0723",
        "title": "【日报】颜浩轩｜2026年7月23日｜组会汇报差异梳理与智能体细节准备",
        "module": "AI + 心理无感知测评项目中，第二次组会汇报准备、智能体细节设计对照与 NK-GeniOS 原型展示材料梳理。",
        "intro": "今日主要在第一周周报和前一次组会汇报的基础上，重新梳理下一次汇报应展示的新增内容，重点把前期访谈需求提炼与智能体具体设计、平台原型和后续代码重构任务衔接起来。",
        "sections": [
            (
                "1. 今天完成了什么？",
                [
                    "对照上一次组会和第一周周报内容，区分已经汇报过的访谈分析、Miro RQ1/RQ2、需求提炼和阶段性报告，避免下一次组会继续重复背景材料。",
                    "重新梳理 `outputs/09_智能体设计` 中的核心文件，重点检查智能体总体设计、对话流程设计、提示词与规则、危机风险识别、心理评估报告模板和 NK-GeniOS 原型制作指南之间的衔接关系。",
                    "把下一次汇报重点收缩为四条主线：智能体具体如何工作、对话流程如何状态化、风险与安全边界如何控制、后台报告和飞书上传如何形成闭环。",
                    "梳理 NK-GeniOS 原型展示思路，将原型目标表述为一条可演示闭环：学生进入智能体、完成边界说明和自然对话、系统结构化追问、完成安全确认、生成后台报告、调用 `report-to-feishu` 并写入飞书记录。",
                    "检查前期文档中可能存在的旧口径，初步发现“人工复核闭环”和当前无感知测评定位之间需要进一步统一，为次日集中修订做准备。",
                ],
            ),
            (
                "2. 今天产出了什么材料？",
                [
                    "形成第二次组会汇报的内容主线草稿，明确本次汇报不再以访谈分析为主，而是以智能体设计细节和可落地原型路径为主。",
                    "整理出汇报需要展示的文件清单，包括智能体总体设计方案、对话流程设计、提示词与规则设计、危机风险识别规则、心理评估报告模板、NK-GeniOS 原型制作指南和后续基础应用重构计划。",
                    "形成报告页面结构建议，包括项目口径变化、智能体模块图、D0-D12 对话流程、报告字段、风险分级、飞书上传闭环和 MVP 后续计划。",
                    "形成待修订问题清单，重点包括人工复核表述、学生端是否展示报告、飞书字段范围、校内资源联系方式和基础应用代码重构边界。",
                ],
            ),
            (
                "3. 今天新增了什么可观察数据或证据？",
                [
                    "明确本次汇报与上次汇报的区别：上次主要说明访谈和 Miro 如何支撑需求，本次需要说明智能体如何具体运行、如何记录状态、如何处理风险、如何生成后台报告。",
                    "从 09 智能体设计文件中提取出可展示的结构化内容，包括系统目标 G1-G5、非目标、模块拆分、D0-D12 流程、提示词八部分结构、风险分级和报告字段。",
                    "确认当前材料已经可以支撑一份图文结合的组会报告，但部分项目口径仍需要在 24 日根据会议讨论和最新要求统一。",
                    "初步将基础应用代码重构定位为后续 MVP 的支撑工作，不在 23 日写成已经完成，只作为下一步计划和汇报预留内容。",
                ],
            ),
            (
                "4. 当前卡点是什么？",
                [
                    "部分旧文档仍沿用人工复核、待复核或专业人员确认等表述，但项目最终形态更接近后台静默报告和飞书自动上传，需要统一口径后再用于正式汇报。",
                    "飞书多维表格字段、上传权限和失败重试机制尚未完全确认，报告中不能把飞书闭环写成已经完成真实联调。",
                    "基础应用代码重构还处于计划阶段，需要先保证不破坏原有可运行页面，再逐步建立工程结构、接口契约和测试材料。",
                ],
            ),
            (
                "5. 明天计划做什么？",
                [
                    "完成第二次组会汇报报告和会议纪要素材整理，重点展示智能体设计细节、无感知测评输出闭环、NK-GeniOS 原型和后续 MVP 计划。",
                    "根据最新项目口径统一修改核心文件，将最终方案明确为后台静默生成报告并自动上传飞书多维表格，取消人工复核闭环表述。",
                    "启动 Mental-LLM 基础工程重构的前两步，优先处理密钥风险和项目目录骨架，建立后续接入智能体、报告生成和飞书工作流的基础结构。",
                ],
            ),
        ],
    },
    {
        "day": "0724",
        "title": "【日报】颜浩轩｜2026年7月24日｜组会汇报、无感知口径修订与基础工程重构",
        "module": "AI + 心理无感知测评项目中，第二次组会汇报、后台自动报告上传方案修订、会议纪要整理和 Mental-LLM 基础工程重构启动。",
        "intro": "今日围绕第二次组会和会后推进展开，重点完成 24 号阶段汇报材料、会议纪要素材、项目口径统一修订，并开始把前期智能体设计转化为 Mental-LLM 项目的工程结构和接口契约。",
        "sections": [
            (
                "1. 今天完成了什么？",
                [
                    "完成 24 号无感知测评项目组会汇报材料，重点展示智能体总体目标、非目标、角色边界、模块拆分、D0-D12 对话流程、提示词规则、风险分级、报告字段和 NK-GeniOS/飞书闭环。",
                    "根据组会内容整理会议纪要素材，将个人项目进展写清楚：当前已完成智能体总体设计、对话流程、安全风险管控、NK 智能 OS 原型设计和后续代码重构计划。",
                    "结合最新项目要求统一核心口径：项目最终不是人工复核系统，而是无感知测评系统；学生端只进行自然对话，不直接查看完整报告；对话结束后系统后台静默生成报告并自动上传飞书多维表格。",
                    "新增并完善“自动报告上传与数据边界设计”相关内容，明确不做人审闭环、不做诊断、不向学生展示完整报告、高关注只给校内心理资源、飞书保存最小必要字段等原则。",
                    "同步修订项目交接说明、访谈需求提炼文档、智能体设计需求清单、智能体总体设计方案、对话流程设计、提示词与规则设计、危机风险识别规则和心理评估报告模板，减少旧版人工复核表述对后续汇报的干扰。",
                    "启动 Mental-LLM 基础工程重构的 Goal 0 和 Goal 1：处理前端 API Key 暴露和本地持久化风险，建立 `.env.example`、`.gitignore`、服务端配置样例、Vite + TypeScript 工具链、目录骨架、接口契约和基础测试文档。",
                ],
            ),
            (
                "2. 今天产出了什么材料？",
                [
                    "《无感知测评汇报报告0724.html》和《无感知测评汇报报告0724.pdf》，用于本次组会展示和阶段性汇报。",
                    "《组会会议纪要素材整理_2026-07-24.md》和《交给Trae生成会议纪要的提示词_2026-07-24.md》，用于生成正式会议纪要。",
                    "《自动报告上传与数据边界设计.md》，用于替代旧的人工复核闭环思路，明确后台报告、飞书上传、学生端资源提示和数据最小化边界。",
                    "《下一次周报展示框架_2026-07-24.md》，用于后续周报展示，重点突出本次与上一次汇报的差异。",
                    "Mental-LLM 工程重构相关材料，包括 `progress/2026-07-24_log.md`、Goal 0/1 测试记录、`package.json`、`tsconfig.json`、`vite.config.ts`、`src/contracts/`、`docs/`、`tests/` 和 `scripts/` 等基础文件。",
                ],
            ),
            (
                "3. 今天新增了什么可观察数据或证据？",
                [
                    "24 号组会报告中已将智能体目标明确为自然收集关键信息、风险优先、降低误导、支持后台自动上传、可实施可测试，同时明确不追求聊天时长、不建立医学诊断模型、不自动通知第三方。",
                    "组会纪要显示，后续最低 MVP 标准被收缩为“能对话、能出报告、能记录必要状态”，这为下一阶段开发提供了更清晰的验收标准。",
                    "后台自动报告上传方案明确飞书字段应包含会话编号、上传时间、主诉摘要、安全关注状态、资源提示、报告版本和上传状态等，同时不建议默认上传完整对话原文和过度身份信息。",
                    "Goal 0 开发记录显示，前端不再持有或发送 API Key，代理侧统一从 `NANKAI_API_KEY` 读取凭据，学生 ID 和姓名只保存在页面内存中。",
                    "Goal 1 测试记录显示，目录骨架、TypeScript 契约、`npm run typecheck`、`npm run build` 和 Python 语法检查已通过，说明基础工程结构可以继续承接后续 MVP 开发。",
                ],
            ),
            (
                "4. 当前卡点是什么？",
                [
                    "飞书多维表格权限和上传流程尚未完全打通，真实写入、字段权限、失败重试和错误记录仍需要后续联调。",
                    "学校心理健康中心联系方式、服务时间和学生端展示文案仍需组内确认，不能在原型中随意填写未确认资源。",
                    "Mental-LLM 目前仍保留旧版 `index.html` 作为实际页面入口，工程骨架已经建立，但视图拆分、接口接入、报告生成和上传链路还需要继续推进。",
                    "已有暴露 Key 的实际轮换需要在外部平台完成，代码侧已经调整，但凭据治理不能只靠本地文件修改完成。",
                ],
            ),
            (
                "5. 明天计划做什么？",
                [
                    "继续推进 Mental-LLM MVP，优先完成对话记录、会话状态、报告草稿和上传状态等最小功能链路。",
                    "根据 24 号组会反馈，继续完善智能体提示词、安全分级和后台报告字段，保证“未询问”和“否定回答”在状态字段中严格区分。",
                    "对接或模拟 `report-to-feishu` 工作流，检查飞书字段映射、上传状态和失败记录是否满足汇报需要。",
                    "准备下一次周报展示材料，增加组会报告截图、流程图、目录结构图、测试记录和 MVP 进展说明，减少纯文字堆叠。",
                ],
            ),
        ],
    },
]


def main():
    paths = []
    for item in REPORTS:
        paths.append(
            build_report(
                item["day"],
                item["title"],
                item["module"],
                item["intro"],
                item["sections"],
            )
        )
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
